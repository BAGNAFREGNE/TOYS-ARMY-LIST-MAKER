import flet as ft
import openpyxl
import os
import warnings
import time
import re

warnings.filterwarnings("ignore")

class ToysArmyApp:
    def __init__(self):
        # --- 1. GESTIONE PERCORSI FILE ---
        self.exe_path = os.path.dirname(os.path.abspath(__file__))
        
        possible_paths = [
            os.path.join(self.exe_path, "assets", "Dati ToysArmy.xlsx"),
            os.path.join(self.exe_path, "Dati ToysArmy.xlsx"),
            "Dati ToysArmy.xlsx"
        ]
        
        self.file_path = "Dati ToysArmy.xlsx"
        for p in possible_paths:
            if os.path.exists(p):
                self.file_path = p
                break
        
        if not os.path.exists(self.file_path):
             if os.path.exists("assets/Dati ToysArmy.xlsx"):
                 self.file_path = "assets/Dati ToysArmy.xlsx"

        # --- 2. VARIABILI ---
        self.list_name = "Nuova_Lista"
        self.selected_units = [] 
        self.compensations = []
        self.notes = [] # NUOVA VARIABILE PER LE NOTE
        self.total_cost = 0.0
        self.alliance = ""
        self.faction = "" 
        
        # VARIABILI BUDGET
        self.budget_limit = 0     
        self.budget_active = False 
        
        self.current_view = "home" 
        self.coming_from = "home" 
        
        self.bg_image = "background.jpg" 

    def main(self, page: ft.Page):
        self.page = page
        self.page.title = "Toys Army"
        self.page.theme_mode = "light" 
        self.page.padding = 0 
        
        self.page.on_back_button = self.handle_back_event
        
        self.show_home()

    def handle_back_event(self, e):
        if self.current_view == "home":
            return False 
        
        elif self.current_view == "saved_lists":
            # Se venivo dalle unità (caricamento interno), torno lì
            if self.coming_from == "units":
                self.show_units()
            else:
                self.show_home()
        
        elif self.current_view == "faction":
            self.show_home()
        
        elif self.current_view == "units":
            self.show_faction()
            
        elif self.current_view == "budget_select": 
            self.show_units()
        
        elif self.current_view == "summary":
            # FIX RIGOROSO: Se vengo da saved_lists, torno a saved_lists
            if self.coming_from == "saved_lists":
                self.show_saved_lists()
            elif self.coming_from == "units":
                self.show_units()
            else:
                self.show_home() 
        return True 

    # --- 3. HELPER EXCEL ---
    def get_excel_data(self, sheet_name):
        try:
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            if sheet_name not in wb.sheetnames:
                return []
            ws = wb[sheet_name]
            data = []
            for row in ws.iter_rows(values_only=True):
                data.append(list(row))
            
            if len(data) > 0:
                return data[1:] 
            return []
        except Exception as e:
            print(f"Errore lettura foglio {sheet_name}: {e}")
            return []

    # --- 4. BOTTONE INTELLIGENTE ---
    def get_smart_button(self, image_name, fallback_text, action, color="blue", size=30):
        image_full_path = os.path.join(self.exe_path, "assets", image_name)
        
        if os.path.exists(image_full_path):
            return ft.Container(
                content=ft.Image(src=image_name, width=size, height=size, fit="contain"), 
                on_click=action,
                padding=5,
                border_radius=5,
                ink=True
            )
        else:
            return ft.ElevatedButton(text=fallback_text, on_click=action, bgcolor=color, color="white")

    # --- 5. LAYOUT ---
    def build_page(self, title_text, main_content, leading_btn=None, footer_content=None, content_alignment=ft.Alignment(0, -1), top_right_widget=None):
        
        header_row = ft.Row([
            leading_btn if leading_btn else ft.Container(width=40),
            
            ft.Text(
                title_text.upper(),
                font_family="Courier New",
                weight="bold", 
                size=20, 
                color="white",
                text_align="center"
            ),
            
            ft.Container(width=40) 
        ], alignment="spaceBetween", vertical_alignment="center") 

        header_container = ft.Container(
            content=header_row,
            padding=ft.padding.only(top=45, left=10, right=10, bottom=10),
            bgcolor="#4D000000", 
        )

        body_container = ft.Container(
            content=main_content, 
            expand=True,          
            padding=5,
            alignment=content_alignment 
        )

        foreground_column = ft.Column([
            header_container,
            body_container
        ], expand=True, spacing=0)

        if footer_content:
            foreground_column.controls.append(footer_content)

        stack_layers = [
            ft.Image(
                src=self.bg_image,
                fit="cover", 
                width=float("inf"),
                height=float("inf"),
                opacity=0.9,
                expand=True
            ),
            foreground_column
        ]

        if top_right_widget:
            stack_layers.append(
                ft.Container(
                    content=top_right_widget,
                    top=45, 
                    right=10,
                )
            )

        return ft.Stack(stack_layers, expand=True)

    # --- 6. HOME ---
    def show_home(self):
        self.current_view = "home"
        self.page.clean()
        self.coming_from = "home"
        
        # Reset totale
        self.list_name = "Nuova_Lista"
        self.selected_units = []
        self.compensations = []
        self.notes = [] # Reset note
        self.total_cost = 0.0
        self.faction = "" 
        self.budget_active = False 
        self.budget_limit = 0
        
        main_content = ft.Column([
                ft.Container(height=40),
                ft.ElevatedButton("ASSE", on_click=lambda _: self.select_alliance("Asse"), height=60, width=200, style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=5))),
                ft.ElevatedButton("ALLEATI", on_click=lambda _: self.select_alliance("Alleati"), height=60, width=200, style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=5))),
                ft.Container(height=20),
                
                ft.ElevatedButton("LISTE SALVATE", on_click=lambda _: self.show_saved_lists(), height=60, width=200, bgcolor="orange", color="white", style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=5))),
                
            ], 
            alignment="center", 
            horizontal_alignment="center", 
            spacing=20,
        )

        self.page.add(self.build_page("TOYS ARMY MENU", main_content, content_alignment=ft.Alignment(0, 0)))
        self.page.update()

    # --- MENU SELEZIONE BUDGET ---
    def show_budget_selector(self):
        self.current_view = "budget_select"
        self.page.clean()
        
        limits = [500, 1000, 1500, 2000, 2500, 3000]
        btn_list = []
        
        btn_list.append(
            ft.ElevatedButton("NESSUN LIMITE", width=250, height=50, bgcolor="green", color="white", 
                              on_click=lambda e: self.set_budget(0, False))
        )
        
        for l in limits:
            btn_list.append(
                ft.ElevatedButton(f"{l} PUNTI", width=250, height=50, 
                                  on_click=lambda e, val=l: self.set_budget(val, True))
            )
            
        col = ft.Column(btn_list, spacing=15, horizontal_alignment="center")
        
        back_btn = self.get_smart_button("Back.png", "<", lambda _: self.show_units())
        
        self.page.add(self.build_page("IMPOSTA LIMITE", col, back_btn, content_alignment=ft.Alignment(0, 0)))
        self.page.update()

    def set_budget(self, limit, active):
        self.budget_limit = limit
        self.budget_active = active
        self.show_units() 

    # --- VISUALIZZA LISTE SALVATE ---
    def show_saved_lists(self):
        self.current_view = "saved_lists"
        self.page.clean()
        
        all_files = [f for f in os.listdir(self.exe_path) if f.endswith(".txt") and f != "requirements.txt"]
        filtered_files = []

        if self.coming_from == "units" and self.faction != "":
            target_faction_line = f"Fazione: {self.faction}"
            for f in all_files:
                try:
                    with open(os.path.join(self.exe_path, f), "r", encoding="utf-8") as file_obj:
                        content = file_obj.read()
                        if target_faction_line in content:
                            filtered_files.append(f)
                except: pass
        else:
            filtered_files = all_files

        list_view = ft.ListView(expand=True, spacing=10, padding=10)
        
        if not filtered_files:
            msg = "Nessuna lista trovata per questa fazione." if self.coming_from == "units" else "Nessuna lista salvata."
            list_view.controls.append(
                ft.Container(
                    content=ft.Text(msg, size=16, weight="bold", font_family="Courier New", text_align="center"),
                    bgcolor="white", padding=10, border_radius=10, alignment=ft.Alignment(0, 0)
                )
            )
        else:
            for filename in filtered_files:
                safe_name = filename.replace(".txt", "")
                
                btn_open = ft.Container(
                    content=ft.Text(safe_name, color="black", size=18, weight="bold", font_family="Courier New"),
                    on_click=lambda e, f=filename: self.load_and_edit_list(f),
                    padding=10,
                    ink=True, 
                    border_radius=5
                )

                row = ft.Row([
                    btn_open,
                    self.get_smart_button("Cancel.png", "X", lambda e, f=filename: self.delete_list(f), color="red", size=25)
                ], alignment="spaceBetween")
                
                container = ft.Container(
                    content=row,
                    padding=5,
                    bgcolor="white",
                    border_radius=5,
                    border=ft.border.all(2, "black")
                )
                list_view.controls.append(container)

        if self.coming_from == "units":
             back_action = lambda _: self.show_units()
        else:
             back_action = lambda _: self.show_home()

        back_btn = self.get_smart_button("Back.png", "<", back_action)
        
        title_page = self.faction.upper() if self.coming_from == "units" else "ARCHIVIO LISTE"
        self.page.add(self.build_page(title_page, list_view, back_btn))
        self.page.update()

    # --- LOGICA CARICAMENTO ---
    def load_and_edit_list(self, filename):
        try:
            full_path = os.path.join(self.exe_path, filename)
            
            temp_units = []
            temp_comp = []
            temp_notes = [] # Temp notes
            temp_cost = 0.0
            temp_faction = ""
            temp_alliance = ""
            temp_name = filename.replace(".txt", "")

            with open(full_path, "r", encoding="utf-8") as f:
                lines = f.readlines()
            
            reading_sconti = False
            reading_note = False # Flag per le note
            
            for line in lines:
                line = line.strip()
                if not line: continue
                
                if line.startswith("Nome Lista:"): pass 
                elif line.startswith("Totale Costo:"): pass
                elif line.startswith("Fazione:"): 
                    temp_faction = line.split(":")[1].strip()
                elif line.startswith("Alleanza:"):
                    temp_alliance = line.split(":")[1].strip()
                    
                elif line.startswith("--- SCONTI ---"): 
                    reading_sconti = True
                    reading_note = False
                elif line.startswith("--- NOTE ---"): # Intercetta sezione note
                    reading_sconti = False
                    reading_note = True
                elif line.startswith("--- ELENCO"): 
                    reading_sconti = False
                    reading_note = False
                else:
                    if reading_sconti:
                        if line.startswith("- "):
                            try:
                                val = float(line.replace("- ", ""))
                                temp_comp.append(val)
                                temp_cost -= val
                            except: pass
                    elif reading_note: # Lettura note
                        if line.startswith("- "):
                            temp_notes.append(line.replace("- ", ""))
                    else:
                        match = re.match(r"^(.*?) x(\d+) \(([\d\.]+)\)$", line)
                        if match:
                            name = match.group(1)
                            qty = int(match.group(2))
                            cost = float(match.group(3))
                            
                            temp_units.append({
                                'id': time.time() + len(temp_units), 
                                'text': line,
                                'cost': cost
                            })
                            temp_cost += cost

            self.selected_units = temp_units
            self.compensations = temp_comp
            self.notes = temp_notes # Assegna note
            self.total_cost = temp_cost
            self.list_name = temp_name
            
            if temp_faction:
                self.faction = temp_faction
                self.alliance = temp_alliance

            # === LOGICA DI NAVIGAZIONE RIGOROSA ===
            
            if self.coming_from == "saved_lists":
                # Se vengo dalle liste salvate, DEVO rimanere in contesto saved_lists
                # per poter tornare indietro correttamente.
                self.coming_from = "saved_lists"
                self.show_summary()
                
            elif self.coming_from == "units":
                # Se ero già dentro un'unità, ricarico l'unità
                self.load_data(self.faction)
            
            else:
                self.show_summary()

            self.page.snack_bar = ft.SnackBar(ft.Text(f"Caricata: {self.list_name}", font_family="Courier New"), bgcolor="green")
            self.page.snack_bar.open = True
            self.page.update()

        except Exception as e:
            self.page.snack_bar = ft.SnackBar(ft.Text(f"Errore: {str(e)}"), bgcolor="red")
            self.page.snack_bar.open = True
            self.page.update()

    def delete_list(self, filename):
        try:
            full_path = os.path.join(self.exe_path, filename)
            os.remove(full_path)
            self.page.snack_bar = ft.SnackBar(ft.Text(f"Eliminata!"), bgcolor="green")
            self.page.snack_bar.open = True
            self.show_saved_lists()
        except Exception as e:
            self.page.snack_bar = ft.SnackBar(ft.Text(f"Errore: {e}"), bgcolor="red")
            self.page.snack_bar.open = True
            self.page.update()

    def select_alliance(self, alliance):
        self.alliance = alliance
        self.show_faction()

    # --- 6. FAZIONE ---
    def show_faction(self):
        self.current_view = "faction"
        self.page.clean()
        factions = ["Unione Sovietica", "Stati Uniti d'America", "Regno Unito"] if self.alliance == "Alleati" else ["Terzo Reich", "Impero Giapponese", "Regno D'Italia"]
        
        buttons = []
        for f in factions:
            buttons.append(ft.ElevatedButton(f.upper(), on_click=lambda e, x=f: self.load_data(x), height=50, width=280, style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=5))))

        main_content = ft.Column(
            buttons, 
            alignment="center", 
            horizontal_alignment="center", 
            spacing=20,
        )

        back_btn = self.get_smart_button("Back.png", "<", lambda _: self.show_home())
        self.page.add(self.build_page(self.alliance, main_content, back_btn, content_alignment=ft.Alignment(0, 0)))
        self.page.update()

    def load_data(self, faction):
        self.faction = faction
        try:
            self.weapons_data = self.get_excel_data(f"Armi {faction}")
            self.roles_data = self.get_excel_data(f"Ruoli {faction}")
            self.tank_data = self.get_excel_data(f"Tank {faction}")
            self.mod_tank_data = self.get_excel_data(f"Mod Tank {faction}")
            
            if self.coming_from != "units" and self.coming_from != "saved_lists" and self.list_name == "Nuova_Lista":
                 self.selected_units = []
                 self.compensations = []
                 self.notes = []
                 self.total_cost = 0.0
                 self.budget_active = False 
            
            self.show_units()
        except Exception as e:
            self.page.snack_bar = ft.SnackBar(ft.Text(f"Errore Dati: {str(e)}"), bgcolor="red")
            self.page.snack_bar.open = True
            self.page.update()

    # --- 7. UNITÀ (CON NOTE) ---
    def show_units(self):
        self.current_view = "units"
        self.page.clean()
        container_list = []
        
        score_label = ft.Text("0", weight="bold", size=14, color="white", font_family="Courier New")
        
        def update_score_display():
            if self.budget_active:
                remaining = self.budget_limit - self.total_cost
                score_label.value = f"RESTANO: {remaining:.0f}"
                score_box.bgcolor = "red"
            else:
                score_label.value = f"TOT: {self.total_cost:.0f}"
                score_box.bgcolor = "green"
            score_label.update()
            score_box.update()

        score_box = ft.Container(
            content=score_label,
            padding=5,
            bgcolor="green", 
            border_radius=5,
            border=ft.border.all(1, "white")
        )
        
        if self.budget_active:
            score_label.value = f"RESTANO: {self.budget_limit - self.total_cost:.0f}"
            score_box.bgcolor = "red"
        else:
            score_label.value = f"TOT: {self.total_cost:.0f}"
            score_box.bgcolor = "green"

        def make_block(label, data, cost_idx, desc_idxs):
            opts = []
            for row in data:
                if row[0] is not None:
                    name = str(row[0])
                    try: cost = row[cost_idx] if row[cost_idx] is not None else 0
                    except: cost = 0
                    opts.append(ft.dropdown.Option(text=f"{name} ({cost})", key=name))

            dd = ft.Dropdown(
                label="Seleziona", 
                options=opts, 
                expand=True,        
                bgcolor="white",
                dense=True,         
                text_size=12,
                content_padding=10,
                menu_height=250 
            )
            
            qty = ft.TextField(
                value="1", label="Qta", width=50, bgcolor="white", 
                text_size=12, content_padding=10, keyboard_type="number"
            )
            
            def add_btn_click(e):
                self.add_item_logic(dd.value, qty.value, data, cost_idx)
                dd.value = None
                update_score_display()
                self.page.snack_bar = ft.SnackBar(ft.Text("Aggiunto!", font_family="Courier New"))
                self.page.snack_bar.open = True
                self.page.update()

            return ft.Container(
                content=ft.Column([
                    ft.Text(label.upper(), weight="bold", size=16, font_family="Courier New"), 
                    ft.Row([dd, qty], alignment="center"), 
                    ft.ElevatedButton("AGGIUNGI", on_click=add_btn_click, style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=2))) 
                ], horizontal_alignment="center"), 
                
                width=320,  
                padding=10,
                border=ft.border.all(1, "black"),
                border_radius=5,
                bgcolor="white",
                alignment=ft.Alignment(0, 0)
            )

        try:
            container_list.append(make_block("Armi", self.weapons_data, 3, [1,2]))
            container_list.append(make_block("Ruoli", self.roles_data, 1, [2]))
            container_list.append(make_block("Tank", self.tank_data, 8, [1]))
            container_list.append(make_block("Modifiche", self.mod_tank_data, 3, [1]))
        except Exception as e:
             self.page.add(ft.Text(f"Errore: {e}", color="red"))

        # --- SEZIONE SCONTI ---
        comp_val = ft.TextField(label="Sconto", width=120, keyboard_type="number", bgcolor="white", text_size=12)
        def apply_comp_click(e):
            try:
                val = float(comp_val.value)
                self.compensations.append(val)
                self.total_cost -= val
                comp_val.value = ""
                update_score_display()
                self.page.snack_bar = ft.SnackBar(ft.Text(f"Sottratto: {val}", font_family="Courier New"))
                self.page.snack_bar.open = True
                self.page.update()
            except: pass

        container_list.append(
            ft.Container(
                content=ft.Column([
                    ft.Text("SCONTI / COMP.", weight="bold", color="red", font_family="Courier New"),
                    ft.Row([comp_val, ft.ElevatedButton("APPLICA", on_click=apply_comp_click, color="white", bgcolor="red")])
                ], horizontal_alignment="center"), 
                width=320, 
                padding=10, border=ft.border.all(1, "red"), border_radius=5, margin=ft.margin.only(top=20), bgcolor="white",
                alignment=ft.Alignment(0, 0)
            )
        )

        # --- SEZIONE NOTE (NUOVA) ---
        note_val = ft.TextField(label="Nota", width=180, bgcolor="white", text_size=12)
        def add_note_click(e):
            if note_val.value:
                self.notes.append(note_val.value)
                note_val.value = ""
                self.page.snack_bar = ft.SnackBar(ft.Text("Nota aggiunta!", font_family="Courier New"))
                self.page.snack_bar.open = True
                self.page.update()

        container_list.append(
            ft.Container(
                content=ft.Column([
                    ft.Text("NOTE", weight="bold", color="blue", font_family="Courier New"),
                    ft.Row([note_val, ft.ElevatedButton("AGGIUNGI", on_click=add_note_click, color="white", bgcolor="blue")])
                ], horizontal_alignment="center"), 
                width=320, 
                padding=10, border=ft.border.all(1, "blue"), border_radius=5, margin=ft.margin.only(top=20), bgcolor="white",
                alignment=ft.Alignment(0, 0)
            )
        )
        
        container_list.append(ft.Container(height=20))

        # --- FOOTER ---
        status_text = f"MODIFICA: {self.list_name}" if self.list_name != "Nuova_Lista" else "NUOVA LISTA"
        status_col = "orange" if self.list_name != "Nuova_Lista" else "green"
        
        lbl_status = ft.Container(
            content=ft.Text(status_text, color=status_col, weight="bold", font_family="Courier New", size=12),
            padding=5,
            bgcolor="#cc000000",
            border_radius=5,
            alignment=ft.Alignment(0, 0)
        )

        def prep_load():
            self.coming_from = "units"
            self.show_saved_lists()

        btn_carica = ft.FloatingActionButton(
            content=ft.Text("CARICA", weight="bold", font_family="Courier New", size=12),
            width=65, height=40, bgcolor="orange",
            on_click=lambda _: prep_load()
        )

        def reset_data_click(e):
            self.list_name = "Nuova_Lista"
            self.selected_units = []
            self.compensations = []
            self.notes = []
            self.total_cost = 0.0
            self.budget_active = False 
            self.show_units() 
            self.page.snack_bar = ft.SnackBar(ft.Text("Lista resettata!", font_family="Courier New"))
            self.page.snack_bar.open = True
            self.page.update()

        btn_nuova = ft.FloatingActionButton(
            content=ft.Text("NUOVA", weight="bold", font_family="Courier New", size=12),
            width=65, height=40, bgcolor="blue",
            on_click=reset_data_click
        )

        btn_set = ft.FloatingActionButton(
            content=ft.Text("SET", weight="bold", font_family="Courier New", size=12),
            width=65, height=40, bgcolor="blue",
            on_click=lambda _: self.show_budget_selector()
        )

        btn_lista = ft.FloatingActionButton(
            content=ft.Text("LISTA", weight="bold", font_family="Courier New", size=12),
            width=65, height=40, bgcolor="green",
            on_click=lambda _: self.go_to_summary_from_units()
        )
        
        footer_content = ft.Container(
            content=ft.Column([
                lbl_status,
                ft.Row([btn_carica, btn_nuova, btn_set, btn_lista], alignment="spaceBetween", width=320)
            ], horizontal_alignment="center", spacing=10),
            padding=ft.padding.only(left=20, right=20, bottom=35, top=10),
            bgcolor=None 
        )

        main_scroll = ft.ListView(
            controls=container_list, 
            expand=True, 
            spacing=15,
            padding=20
        )
        
        back_btn = self.get_smart_button("Back.png", "<", lambda _: self.show_faction())
        
        self.page.add(self.build_page(self.faction, main_scroll, back_btn, footer_content=footer_content, top_right_widget=score_box))
        self.page.update()

    def go_to_summary_from_units(self):
        self.coming_from = "units"
        self.show_summary()

    def add_item_logic(self, name, qty_str, data, cost_idx):
        if not name: return
        try: q = int(qty_str)
        except: q = 1
        
        cost = 0
        for row in data:
            if str(row[0]) == name:
                try: cost = float(row[cost_idx])
                except: cost = 0
                break
        
        tot = cost * q
        self.selected_units.append({'id': time.time(), 'text': f"{name} x{q} ({tot})", 'cost': tot})
        self.total_cost += tot

    # --- 8. RIEPILOGO ---
    def show_summary(self):
        self.current_view = "summary"
        self.page.clean()
        
        list_view = ft.ListView(expand=True, spacing=10, padding=10)

        def refresh_list_view():
            list_view.controls.clear()
            
            name_field = ft.TextField(label="NOME LISTA", value=self.list_name, on_change=lambda e: setattr(self, 'list_name', e.control.value), bgcolor="white", width=320, text_style=ft.TextStyle(font_family="Courier New"))
            list_view.controls.append(ft.Container(content=name_field, alignment=ft.Alignment(0, 0))) 
            list_view.controls.append(ft.Divider(color="black"))

            items_col = ft.Column()
            
            if not self.selected_units and not self.compensations and not self.notes:
                items_col.controls.append(ft.Text("LISTA VUOTA", font_family="Courier New"))
            
            # Unità
            for u in self.selected_units:
                row = ft.Row([
                    ft.Text(u['text'], expand=True, font_family="Courier New", weight="bold"),
                    self.get_smart_button("Cancel.png", "X", lambda e, x=u: remove_unit(x), color="red", size=25)
                ], alignment="spaceBetween") 
                items_col.controls.append(row)

            # Sconti
            if self.compensations:
                items_col.controls.append(ft.Divider(color="black"))
                items_col.controls.append(ft.Text("SCONTI:", italic=True, color="red", font_family="Courier New"))
                for c in self.compensations:
                    row = ft.Row([
                        ft.Text(f"- {c}", color="red", expand=True, font_family="Courier New"),
                        self.get_smart_button("Cancel.png", "X", lambda e, x=c: remove_comp(x), color="red", size=25)
                    ], alignment="spaceBetween") 
                    items_col.controls.append(row)

            # Note (Visualizzazione)
            if self.notes:
                items_col.controls.append(ft.Divider(color="black"))
                items_col.controls.append(ft.Text("NOTE:", italic=True, color="blue", font_family="Courier New"))
                for n in self.notes:
                    row = ft.Row([
                        ft.Text(f"- {n}", color="blue", expand=True, font_family="Courier New"),
                        self.get_smart_button("Cancel.png", "X", lambda e, x=n: remove_note(x), color="red", size=25)
                    ], alignment="spaceBetween") 
                    items_col.controls.append(row)
            
            list_view.controls.append(ft.Container(content=items_col, border=ft.border.all(2, "black"), border_radius=5, padding=5, bgcolor="white"))
            list_view.controls.append(ft.Divider(color="black"))
            
            # Totale
            if self.budget_active:
                remaining = self.budget_limit - self.total_cost
                total_text = f"RESTANO: {remaining:.0f} (Su {self.budget_limit})"
                txt_col = "red"
            else:
                total_text = f"TOTALE: {self.total_cost:.2f}"
                txt_col = "black"

            total_label = ft.Text(total_text, size=25, weight="bold", color=txt_col, font_family="Courier New", bgcolor="white")
            list_view.controls.append(ft.Container(content=total_label, alignment=ft.Alignment(0, 0)))
            
            save_btn = ft.ElevatedButton("SALVA LISTA", on_click=save_file, bgcolor="green", color="white", style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=2)))
            list_view.controls.append(ft.Container(content=save_btn, alignment=ft.Alignment(0, 0)))
            
            if self.coming_from != "units" and self.faction:
                def go_to_edit_mode(e):
                    self.coming_from = "units" 
                    self.load_data(self.faction) 
                
                edit_btn = ft.ElevatedButton("AGGIUNGI PEZZI", on_click=go_to_edit_mode, bgcolor="blue", color="white", style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=2)))
                list_view.controls.append(ft.Container(content=edit_btn, alignment=ft.Alignment(0, 0), margin=ft.margin.only(top=10)))

            self.page.update()

        def remove_unit(unit_obj):
            self.selected_units.remove(unit_obj)
            self.total_cost -= unit_obj['cost']
            refresh_list_view()

        def remove_comp(comp_val):
            self.compensations.remove(comp_val)
            self.total_cost += comp_val
            refresh_list_view()

        def remove_note(note_val):
            self.notes.remove(note_val)
            refresh_list_view()

        def save_file(e):
            try:
                safe_name = "".join([c for c in self.list_name if c.isalnum() or c in (' ', '_')]).strip()
                if not safe_name: safe_name = "Lista"
                filename = f"{safe_name}.txt"
                full_path = os.path.join(self.exe_path, filename)
                
                with open(full_path, "w", encoding="utf-8") as f:
                    f.write(f"Nome Lista: {self.list_name}\n")
                    f.write(f"Fazione: {self.faction}\n") 
                    f.write(f"Alleanza: {self.alliance}\n")
                    f.write(f"Totale Costo: {self.total_cost}\n\n")
                    f.write("--- ELENCO UNITÀ ---\n")
                    for u in self.selected_units: f.write(u['text'] + "\n")
                    if self.compensations:
                        f.write("\n--- SCONTI ---\n")
                        for c in self.compensations: f.write(f"- {c}\n")
                    if self.notes:
                        f.write("\n--- NOTE ---\n")
                        for n in self.notes: f.write(f"- {n}\n")
                        
                self.page.snack_bar = ft.SnackBar(ft.Text(f"SALVATO: {filename}", font_family="Courier New"), bgcolor="green", duration=4000)
                self.page.snack_bar.open = True
                self.page.update()
            except Exception as ex:
                self.page.snack_bar = ft.SnackBar(ft.Text(f"Errore: {ex}"), bgcolor="red")
                self.page.snack_bar.open = True
                self.page.update()

        refresh_list_view()

        if self.coming_from == "saved_lists":
            back_action = lambda _: self.show_saved_lists()
        else:
            back_action = lambda _: self.show_units()
            
        back_btn = self.get_smart_button("Back.png", "<", back_action)
        self.page.add(self.build_page("RIEPILOGO", list_view, back_btn))
        self.page.update()

if __name__ == "__main__":
    app = ToysArmyApp()
    assets_folder = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
    if os.path.exists(assets_folder):
        ft.app(target=app.main, assets_dir="assets")
    else:
        ft.app(target=app.main)
